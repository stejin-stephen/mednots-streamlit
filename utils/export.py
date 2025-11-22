from datetime import datetime
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from fhir.resources.documentreference import DocumentReference
from fhir.resources.composition import Composition, CompositionSection
from fhir.resources.narrative import Narrative
import json

def export_to_pdf(soap_data: dict, vitals_data: dict, findings_data: dict, key_info: dict) -> BytesIO:
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []
    
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#1f77b4'),
        spaceAfter=12
    )
    
    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading2'],
        fontSize=14,
        textColor=colors.HexColor('#2c3e50'),
        spaceAfter=8,
        spaceBefore=12
    )
    
    story.append(Paragraph("Medical Note Summary", title_style))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
    story.append(Spacer(1, 0.3*inch))
    
    if key_info:
        story.append(Paragraph("Key Information", heading_style))
        for key, value in key_info.items():
            story.append(Paragraph(f"<b>{key.replace('_', ' ').title()}:</b> {value}", styles['Normal']))
        story.append(Spacer(1, 0.2*inch))
    
    if vitals_data and vitals_data.get("key_vitals"):
        story.append(Paragraph("Vital Signs", heading_style))
        vitals = vitals_data["key_vitals"]
        vital_items = []
        for key, value in vitals.items():
            if value and value != "N/A":
                vital_items.append([key.replace('_', ' ').title(), value])
        
        if vital_items:
            t = Table(vital_items, colWidths=[2.5*inch, 3*inch])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f0f0f0')),
                ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey)
            ]))
            story.append(t)
        story.append(Spacer(1, 0.2*inch))
    
    if soap_data and soap_data.get("soap_note"):
        soap = soap_data["soap_note"]
        story.append(Paragraph("SOAP Note", heading_style))
        
        for section, label in [("subjective", "Subjective"), ("objective", "Objective"), 
                               ("assessment", "Assessment"), ("plan", "Plan")]:
            content = soap.get(section, "N/A")
            story.append(Paragraph(f"<b>{label}:</b>", styles['Normal']))
            story.append(Paragraph(content, styles['Normal']))
            story.append(Spacer(1, 0.1*inch))
    
    if findings_data and findings_data.get("exam_findings"):
        story.append(Paragraph("Examination Findings", heading_style))
        for i, finding in enumerate(findings_data["exam_findings"], 1):
            story.append(Paragraph(f"{i}. {finding}", styles['Normal']))
    
    doc.build(story)
    buffer.seek(0)
    return buffer

def export_to_word(soap_data: dict, vitals_data: dict, findings_data: dict, key_info: dict) -> BytesIO:
    doc = Document()
    
    title = doc.add_heading('Medical Note Summary', 0)
    title.runs[0].font.color.rgb = (31, 119, 180)
    
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    
    if key_info:
        doc.add_heading('Key Information', level=1)
        for key, value in key_info.items():
            p = doc.add_paragraph()
            p.add_run(f"{key.replace('_', ' ').title()}: ").bold = True
            p.add_run(value)
    
    if vitals_data and vitals_data.get("key_vitals"):
        doc.add_heading('Vital Signs', level=1)
        vitals = vitals_data["key_vitals"]
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Light Grid Accent 1'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Vital'
        hdr_cells[1].text = 'Value'
        
        for key, value in vitals.items():
            if value and value != "N/A":
                row_cells = table.add_row().cells
                row_cells[0].text = key.replace('_', ' ').title()
                row_cells[1].text = value
    
    if soap_data and soap_data.get("soap_note"):
        soap = soap_data["soap_note"]
        doc.add_heading('SOAP Note', level=1)
        
        for section, label in [("subjective", "Subjective"), ("objective", "Objective"), 
                               ("assessment", "Assessment"), ("plan", "Plan")]:
            content = soap.get(section, "N/A")
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            doc.add_paragraph(content)
    
    if findings_data and findings_data.get("exam_findings"):
        doc.add_heading('Examination Findings', level=1)
        for i, finding in enumerate(findings_data["exam_findings"], 1):
            doc.add_paragraph(f"{i}. {finding}", style='List Number')
    
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def export_to_fhir(soap_data: dict, vitals_data: dict, findings_data: dict, key_info: dict) -> str:
    composition_sections = []
    
    if soap_data and soap_data.get("soap_note"):
        soap = soap_data["soap_note"]
        for code, title in [("subjective", "Subjective"), ("objective", "Objective"), 
                           ("assessment", "Assessment"), ("plan", "Plan")]:
            section = CompositionSection(
                title=title,
                text=Narrative(
                    status="generated",
                    div=f"<div xmlns='http://www.w3.org/1999/xhtml'><p>{soap.get(code, 'N/A')}</p></div>"
                )
            )
            composition_sections.append(section)
    
    if vitals_data and vitals_data.get("key_vitals"):
        vitals_html = "<div xmlns='http://www.w3.org/1999/xhtml'><ul>"
        for key, value in vitals_data["key_vitals"].items():
            if value and value != "N/A":
                vitals_html += f"<li><b>{key.replace('_', ' ').title()}:</b> {value}</li>"
        vitals_html += "</ul></div>"
        
        vitals_section = CompositionSection(
            title="Vital Signs",
            text=Narrative(status="generated", div=vitals_html)
        )
        composition_sections.append(vitals_section)
    
    if findings_data and findings_data.get("exam_findings"):
        findings_html = "<div xmlns='http://www.w3.org/1999/xhtml'><ul>"
        for finding in findings_data["exam_findings"]:
            findings_html += f"<li>{finding}</li>"
        findings_html += "</ul></div>"
        
        findings_section = CompositionSection(
            title="Examination Findings",
            text=Narrative(status="generated", div=findings_html)
        )
        composition_sections.append(findings_section)
    
    composition = Composition(
        status="final",
        type={
            "coding": [{
                "system": "http://loinc.org",
                "code": "11488-4",
                "display": "Consult note"
            }]
        },
        date=datetime.now().isoformat(),
        title="Clinical Note Summary",
        section=composition_sections
    )
    
    return composition.json(indent=2)
